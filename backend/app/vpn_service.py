from __future__ import annotations

import csv
import io
import os
import re
import shlex
import socket
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import paramiko
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlmodel import select

from .database import get_session
from .models.portal import DiagnosticExportTable, VpnExportTask, VpnSite
BACKEND_ROOT = Path(__file__).resolve().parent.parent
EXPORT_DIR = BACKEND_ROOT / "uploads" / "vpn_exports"
EXPORT_DIR.mkdir(parents=True, exist_ok=True)


class VpnConfigurationError(RuntimeError):
    pass


KEYCHAIN_SERVICES = {
    "VPN_JUMP_PASSWORD": "jd-energy-vpn-jump",
    "VPN_SITE_PASSWORD": "jd-energy-vpn-site",
    "TDENGINE_PASSWORD": "jd-energy-tdengine",
}

SETTING_DEFAULTS = {
    "VPN_JUMP_HOST": "8.216.40.206",
    "VPN_JUMP_USERNAME": "root",
    "VPN_SITE_USERNAME": "root",
    "TDENGINE_USERNAME": "root",
}


def keychain_password(service_name: str) -> str:
    if sys.platform != "darwin":
        return ""
    try:
        result = subprocess.run(
            ["/usr/bin/security", "find-generic-password", "-a", "root", "-s", service_name, "-w"],
            capture_output=True,
            check=False,
            text=True,
            timeout=5,
        )
    except (OSError, subprocess.SubprocessError):
        return ""
    return result.stdout.strip() if result.returncode == 0 else ""


def required_setting(name: str) -> str:
    value = os.getenv(name, "").strip()
    if not value and name in SETTING_DEFAULTS:
        value = SETTING_DEFAULTS[name]
    if not value and name in KEYCHAIN_SERVICES:
        value = keychain_password(KEYCHAIN_SERVICES[name])
    if not value:
        raise VpnConfigurationError(f"Server setting {name} is not configured in the environment or Keychain")
    return value


def ssh_settings() -> Dict[str, object]:
    return {
        "jump_host": required_setting("VPN_JUMP_HOST"),
        "jump_port": int(os.getenv("VPN_JUMP_PORT", "22")),
        "jump_username": required_setting("VPN_JUMP_USERNAME"),
        "jump_password": required_setting("VPN_JUMP_PASSWORD"),
        "site_port": int(os.getenv("VPN_SITE_PORT", "22")),
        "site_username": required_setting("VPN_SITE_USERNAME"),
        "site_password": required_setting("VPN_SITE_PASSWORD"),
        "tdengine_username": required_setting("TDENGINE_USERNAME"),
        "tdengine_password": required_setting("TDENGINE_PASSWORD"),
        "strict_host_key": os.getenv("VPN_SSH_STRICT_HOST_KEY", "true").lower() not in {"0", "false", "no"},
        "known_hosts": os.getenv("VPN_SSH_KNOWN_HOSTS", "").strip(),
    }


def new_ssh_client(settings: Dict[str, object]) -> paramiko.SSHClient:
    client = paramiko.SSHClient()
    known_hosts = str(settings["known_hosts"])
    if known_hosts:
        client.load_host_keys(known_hosts)
    else:
        client.load_system_host_keys()
    if bool(settings["strict_host_key"]):
        client.set_missing_host_key_policy(paramiko.RejectPolicy())
    else:
        client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    return client


def connect_jump_host() -> Tuple[paramiko.SSHClient, Dict[str, object]]:
    settings = ssh_settings()
    client = new_ssh_client(settings)
    client.connect(
        hostname=str(settings["jump_host"]),
        port=int(settings["jump_port"]),
        username=str(settings["jump_username"]),
        password=str(settings["jump_password"]),
        timeout=15,
        banner_timeout=15,
        auth_timeout=15,
        look_for_keys=False,
        allow_agent=False,
    )
    return client, settings


def connect_site(site: VpnSite) -> Tuple[paramiko.SSHClient, paramiko.SSHClient]:
    jump_client, settings = connect_jump_host()
    transport = jump_client.get_transport()
    if transport is None or not transport.is_active():
        jump_client.close()
        raise RuntimeError("Jump host SSH transport is unavailable")
    channel = transport.open_channel(
        "direct-tcpip",
        (site.vpn_ip, int(settings["site_port"])),
        ("127.0.0.1", 0),
    )
    site_client = new_ssh_client(settings)
    try:
        site_client.connect(
            hostname=site.vpn_ip,
            port=int(settings["site_port"]),
            username=str(settings["site_username"]),
            password=str(settings["site_password"]),
            sock=channel,
            timeout=15,
            banner_timeout=15,
            auth_timeout=15,
            look_for_keys=False,
            allow_agent=False,
        )
    except Exception:
        channel.close()
        jump_client.close()
        raise
    return jump_client, site_client


def sql_for_table(table: DiagnosticExportTable, eblock_id: int, start_time: datetime, end_time: datetime) -> str:
    start = start_time.strftime("%Y-%m-%d %H:%M:%S")
    end = end_time.strftime("%Y-%m-%d %H:%M:%S")
    conditions = []
    if table.extra_where:
        conditions.append(table.extra_where)
    if table.has_eblock_id:
        conditions.append(f"eBlock_id = {int(eblock_id)}")
    conditions.extend((f"time >= '{start}'", f"time <= '{end}'"))
    return f"SELECT * FROM {table.table_name} WHERE {' AND '.join(conditions)};"


def build_export_command(
    tables: Iterable[DiagnosticExportTable],
    eblock_id: int,
    start_time: datetime,
    end_time: datetime,
    remote_dir: str,
    tdengine_username: str,
    tdengine_password: str,
) -> Tuple[str, Dict[str, str]]:
    remote_files: Dict[str, str] = {}
    commands = [f"mkdir -p {shlex.quote(remote_dir)}"]
    for table in tables:
        sql = sql_for_table(table, eblock_id, start_time, end_time)
        remote_path = f"{remote_dir}/{table.sheet_name}.csv"
        remote_files[table.sheet_name] = remote_path
        taos_command = " ".join(
            [
                "taos",
                f"-u{shlex.quote(tdengine_username)}",
                f"-p{shlex.quote(tdengine_password)}",
                "-s",
                shlex.quote(sql),
                ">",
                shlex.quote(remote_path),
            ]
        )
        commands.append(taos_command)
    return "set -e; " + " && ".join(commands), remote_files


def normalize_csv_rows(raw: str) -> List[List[str]]:
    text = raw.lstrip("\ufeff").strip()
    if not text:
        return []
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t|")
        return list(csv.reader(io.StringIO(text), dialect))
    except csv.Error:
        return [[part.strip() for part in line.split()] for line in text.splitlines() if line.strip()]


def safe_export_name(site_name: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9._-]+", "-", site_name.strip()).strip(".-")
    return normalized[:80] or "vpn-site"


def excel_safe_value(value: str) -> str:
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def build_workbook(csv_files: Dict[str, Path], destination: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="E2E8F0")
    for table, csv_path in csv_files.items():
        sheet = workbook.create_sheet(title=table)
        raw = csv_path.read_text(encoding="utf-8", errors="replace")
        rows = normalize_csv_rows(raw)
        if not rows:
            rows = [["No data"]]
        for row_index, row in enumerate(rows, start=1):
            sheet.append([excel_safe_value(value) for value in row])
            if row_index == 1:
                for cell in sheet[row_index]:
                    cell.font = Font(bold=True)
                    cell.fill = header_fill
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    workbook.save(destination)


def run_export_task(task_id: int) -> None:
    jump_client: Optional[paramiko.SSHClient] = None
    site_client: Optional[paramiko.SSHClient] = None
    remote_dir = f"/tmp/jd-vpn-export-{task_id}"
    try:
        with get_session() as session:
            task = session.get(VpnExportTask, task_id)
            if task is None:
                return
            site = session.get(VpnSite, task.site_id)
            if site is None:
                raise RuntimeError("VPN site no longer exists")
            task_data = task.model_dump()
            site_data = site.model_dump()
            configured_tables = session.exec(
                select(DiagnosticExportTable).where(
                    DiagnosticExportTable.table_name.in_(task.selected_tables)
                )
            ).all()

        configured_by_name = {item.table_name: item for item in configured_tables}
        export_tables = [configured_by_name[name] for name in task_data["selected_tables"] if name in configured_by_name]
        if len(export_tables) != len(task_data["selected_tables"]):
            raise RuntimeError("One or more diagnostic table configurations no longer exist")

        task = VpnExportTask(**task_data)
        site = VpnSite(**site_data)
        settings = ssh_settings()
        jump_client, site_client = connect_site(site)
        command, remote_files = build_export_command(
            export_tables,
            task.eblock_id,
            task.start_time,
            task.end_time,
            remote_dir,
            str(settings["tdengine_username"]),
            str(settings["tdengine_password"]),
        )
        _, stdout, stderr = site_client.exec_command(command, timeout=1800)
        exit_code = stdout.channel.recv_exit_status()
        error_output = stderr.read().decode("utf-8", errors="replace").strip()
        if exit_code != 0:
            raise RuntimeError(error_output or f"Remote export command exited with code {exit_code}")

        stamp = task.created_at.strftime("%Y%m%d-%H%M%S")
        safe_site_name = safe_export_name(task.site_name)
        archive_name = f"{safe_site_name}-{stamp}.zip"
        archive_path = EXPORT_DIR / archive_name
        with tempfile.TemporaryDirectory(prefix=f"vpn-export-{task_id}-") as temp_dir_name:
            temp_dir = Path(temp_dir_name)
            local_csv_files: Dict[str, Path] = {}
            sftp = site_client.open_sftp()
            try:
                for table, remote_path in remote_files.items():
                    local_path = temp_dir / f"{table}.csv"
                    sftp.get(remote_path, str(local_path))
                    local_csv_files[table] = local_path
            finally:
                sftp.close()
            workbook_path = temp_dir / f"{safe_site_name}-{stamp}.xlsx"
            build_workbook(local_csv_files, workbook_path)
            with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                for table, local_path in local_csv_files.items():
                    archive.write(local_path, arcname=f"csv/{table}.csv")
                archive.write(workbook_path, arcname=workbook_path.name)

        site_client.exec_command(f"rm -rf {shlex.quote(remote_dir)}")
        with get_session() as session:
            stored_task = session.get(VpnExportTask, task_id)
            if stored_task is not None:
                stored_task.status = "completed"
                stored_task.file_name = archive_name
                stored_task.file_path = str(archive_path)
                stored_task.completed_at = datetime.utcnow()
                session.add(stored_task)
                session.commit()
    except Exception as exc:
        with get_session() as session:
            stored_task = session.get(VpnExportTask, task_id)
            if stored_task is not None:
                stored_task.status = "failed"
                stored_task.error_message = str(exc)[:1000]
                stored_task.completed_at = datetime.utcnow()
                session.add(stored_task)
                session.commit()
    finally:
        if site_client is not None:
            site_client.close()
        if jump_client is not None:
            jump_client.close()


def open_terminal_channel(site: Optional[VpnSite], cols: int = 120, rows: int = 32):
    if site is None:
        jump_client, _ = connect_jump_host()
        channel = jump_client.invoke_shell(term="xterm-256color", width=cols, height=rows)
        channel.settimeout(0.2)
        return jump_client, None, channel
    jump_client, site_client = connect_site(site)
    channel = site_client.invoke_shell(term="xterm-256color", width=cols, height=rows)
    channel.settimeout(0.2)
    return jump_client, site_client, channel


def receive_channel(channel: paramiko.Channel) -> bytes:
    try:
        return channel.recv(32768)
    except socket.timeout:
        return b""
